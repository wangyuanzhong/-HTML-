"""
生成与网页矩阵「同版面」的两份通用表格。

输出：data/matrix-objective.xlsx、data/matrix-subjective.xlsx

工作表：
  「网页矩阵」——与浏览器里表格一致：
    • 第 1 行：A 列左上角文案 (= cornerLabel)，B 起为各产品列标题
    • 第 2 行：A 放空；B 起为内部列 ID（如 p1、p2、p3）
    • 第 3 行起：A 列为左侧对比项文案；B 起为格子摘要（与网页 tbody 对齐）
  「详情」——（可选）与「网页矩阵」表头及 A 列对齐；自第一个对比项起每一项占连续两行：
      第 1 行（文）：B 起为弹窗内说明文字（纯文本或多段用空行分段；以 < 开头则视为 HTML）
      第 2 行（图）：B 起为图片路径（相对 HTML 或 file 可访问的路径；可空）
    A 列第 2 行可为「（图路径）」提示列，可清空。

仍会兼容旧的 Rows/Columns/Meta/Cells 工作簿；新格式优先解析。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"

CORNER = r"维度 \ 对象"

ROWS = [
    ("dimension_1", "维度一"),
    ("dimension_2", "维度二"),
    ("dimension_3", "维度三"),
    ("dimension_4", "维度四"),
]

COLS = [
    ("item_a", "对象 A"),
    ("item_b", "对象 B"),
    ("item_c", "对象 C"),
]


def _cell_map(
    seq: list[tuple[str, str, str, str, str]],
) -> dict[tuple[str, str], tuple[str, str, str]]:
    """(row_id, col_id) -> (summary, detail_plain_text, image_path)"""
    d: dict[tuple[str, str], tuple[str, str, str]] = {}
    for rid, cid, summary, text, image_path in seq:
        d[(rid, cid)] = (summary, text, image_path)
    return d


# 摘要 + 详情文 + 图路径（请将图片放到与 HTML 相对路径下，或改写为实际 URL）
_RAW: list[tuple[str, str, str, str, str]] = [
    ("dimension_1", "item_a", "内容 A1", "这里填写对象 A 在维度一的说明。", ""),
    ("dimension_1", "item_b", "内容 B1", "这里填写对象 B 在维度一的说明。", ""),
    ("dimension_1", "item_c", "内容 C1", "这里填写对象 C 在维度一的说明。", ""),
    ("dimension_2", "item_a", "内容 A2", "这里填写对象 A 在维度二的说明。", ""),
    ("dimension_2", "item_b", "内容 B2", "这里填写对象 B 在维度二的说明。", ""),
    ("dimension_2", "item_c", "内容 C2", "这里填写对象 C 在维度二的说明。", ""),
    ("dimension_3", "item_a", "内容 A3", "这里填写对象 A 在维度三的说明。", ""),
    ("dimension_3", "item_b", "内容 B3", "这里填写对象 B 在维度三的说明。", ""),
    ("dimension_3", "item_c", "内容 C3", "这里填写对象 C 在维度三的说明。", ""),
    ("dimension_4", "item_a", "内容 A4", "这里填写对象 A 在维度四的说明。", ""),
    ("dimension_4", "item_b", "内容 B4", "这里填写对象 B 在维度四的说明。", ""),
    ("dimension_4", "item_c", "内容 C4", "这里填写对象 C 在维度四的说明。", ""),
]

TECH_CELLS = _cell_map(_RAW)
SUBJECTIVE_CELLS = _cell_map(list(_RAW))


def _fill_web_grid(ws, cells: dict[tuple[str, str], tuple[str, str, str]]) -> int:
    """写入「网页矩阵」表头 + 正文，返回正文起始行（1-based Excel）"""
    ws["A1"] = CORNER
    for ci, (_, lab) in enumerate(COLS, start=2):
        ws.cell(row=1, column=ci, value=lab)
    ws["A2"] = ""
    for ci, (cid, _) in enumerate(COLS, start=2):
        ws.cell(row=2, column=ci, value=cid)
    r_out = 3
    for rid, rlab in ROWS:
        ws.cell(row=r_out, column=1, value=rlab)
        for ci, (cid, _) in enumerate(COLS, start=2):
            summary, _, _ = cells.get((rid, cid), ("", "", ""))
            ws.cell(row=r_out, column=ci, value=summary)
        r_out += 1
    return 3


def _fill_detail_sheet(
    ws, data_start_row: int, cells: dict[tuple[str, str], tuple[str, str, str]]
) -> None:
    """与网页矩阵表头 + A 列对比项对齐；每个对比项两行：文 / 图片路径"""
    ws["A1"] = "明细：每项两行（文 → 图路径），与网页矩阵同列对齐"
    for ci, (_, lab) in enumerate(COLS, start=2):
        ws.cell(row=1, column=ci, value=lab)
    ws["A2"] = ""
    for ci, (cid, _) in enumerate(COLS, start=2):
        ws.cell(row=2, column=ci, value=cid)
    r_out = data_start_row
    for rid, rlab in ROWS:
        ws.cell(row=r_out, column=1, value=rlab)
        for ci, (cid, _) in enumerate(COLS, start=2):
            _, text, _ = cells.get((rid, cid), ("", "", ""))
            ws.cell(row=r_out, column=ci, value=text)
        r_out += 1
        ws.cell(row=r_out, column=1, value="（图路径）")
        for ci, (cid, _) in enumerate(COLS, start=2):
            _, _, img_path = cells.get((rid, cid), ("", "", ""))
            ws.cell(row=r_out, column=ci, value=img_path or "")
        r_out += 1


def _autowidth(ws) -> None:
    for col in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 120), max_col=ws.max_column):
        letter = get_column_letter(col[0].column)
        w = 10
        for c in col:
            if c.value is None:
                continue
            w = max(w, min(60, len(str(c.value))))
        ws.column_dimensions[letter].width = w + 1


def write_workbook(path: Path, cells: dict[tuple[str, str], tuple[str, str, str]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws_grid = wb.create_sheet("网页矩阵")
    ws_det = wb.create_sheet("详情")

    ds = _fill_web_grid(ws_grid, cells)
    _fill_detail_sheet(ws_det, ds, cells)

    for ws in (ws_grid, ws_det):
        _autowidth(ws)

    wb.save(path)


def main() -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    write_workbook(DATA / "matrix-objective.xlsx", TECH_CELLS)
    write_workbook(DATA / "matrix-subjective.xlsx", SUBJECTIVE_CELLS)
    print("written:", DATA / "matrix-objective.xlsx")
    print("written:", DATA / "matrix-subjective.xlsx")


if __name__ == "__main__":
    main()
